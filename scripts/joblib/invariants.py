#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""数据层判据（assert 步骤的内置不变量）。

只依赖标准库：csv / json / jsonl 一定支持；
.xlsx 需要 openpyxl，缺失时归为 env 类失败并给出安装提示（不静默跳过）。
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .core import JobError


class DataError(JobError):
    """数据/断言类失败（对应失效分类 data）。"""


class EnvError(JobError):
    """环境类失败（缺依赖、缺文件、没权限）。"""


# ---------------------------------------------------------------- 读记录

def load_records(path, base: Path) -> list:
    p = Path(str(path))
    if not p.is_absolute():
        p = Path(base) / p
    if not p.is_file():
        raise EnvError(f"数据文件不存在：{p}")
    suffix = p.suffix.lower()
    if suffix == ".csv":
        with p.open(newline="", encoding="utf-8-sig") as fh:
            return [dict(r) for r in csv.DictReader(fh)]
    if suffix == ".jsonl":
        rows = []
        for line in p.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line:
                rows.append(json.loads(line))
        return rows
    if suffix == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            for k in ("rows", "records", "data", "items"):
                if isinstance(data.get(k), list):
                    return data[k]
            raise DataError(f"{p} 是对象但没有 rows/records/data/items 数组")
        if isinstance(data, list):
            return data
        raise DataError(f"{p} 不是记录数组")
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:
            raise EnvError(
                f"读 {p.name} 需要 openpyxl（当前 python3 未安装）。"
                f"安装：python3 -m pip install openpyxl；"
                f"或让上游 tool 步骤先把表转成 csv。") from exc
        wb = load_workbook(p, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        return [dict(zip(header, r)) for r in rows[1:]]
    raise DataError(f"不认识的数据格式：{p.name}（支持 csv / json / jsonl / xlsx）")


def _col(row: dict, name: str):
    if name in row:
        return row[name]
    for k in row:
        if str(k).strip() == str(name).strip():
            return row[k]
    raise DataError(f"记录里没有列「{name}」（实际列：{', '.join(map(str, row.keys()))}）")


def _num(v) -> float:
    if v is None or v == "":
        raise DataError("数值列为空")
    s = str(v).replace(",", "").replace("¥", "").replace("￥", "").strip()
    try:
        return float(s)
    except ValueError:
        raise DataError(f"不是数字：{v!r}")


# ---------------------------------------------------------------- 内置不变量

def inv_file_exists(args, base, ctx):
    pats = args.get("paths") or args.get("path")
    from .core import expand_globs
    hits = expand_globs(pats, base)
    if not hits:
        raise DataError(f"期望存在的产物不存在：{pats}")
    return f"{len(hits)} 个文件存在"


def inv_not_empty(args, base, ctx):
    pats = args.get("paths") or args.get("path")
    from .core import expand_globs
    hits = expand_globs(pats, base)
    if not hits:
        raise DataError(f"产物为空（没有匹配到文件）：{pats}")
    empty = [h for h in hits if (base / h).is_file() and (base / h).stat().st_size == 0]
    if empty:
        raise DataError(f"文件为空：{', '.join(empty)}")
    return f"{len(hits)} 个文件非空"


def inv_row_count_between(args, base, ctx):
    rows = load_records(args["path"], base)
    lo, hi = args.get("min", 0), args.get("max")
    if hi is not None and not (lo <= len(rows) <= hi):
        raise DataError(f"{args['path']} 行数 {len(rows)} 不在 [{lo}, {hi}]")
    if hi is None and len(rows) < lo:
        raise DataError(f"{args['path']} 行数 {len(rows)} < {lo}")
    return f"{args['path']} 行数 {len(rows)}"


def inv_columns_present(args, base, ctx):
    rows = load_records(args["path"], base)
    if not rows:
        raise DataError(f"{args['path']} 没有数据，无法校验列")
    cols = set(map(str, rows[0].keys()))
    missing = [c for c in args.get("columns", []) if c not in cols]
    if missing:
        raise DataError(f"{args['path']} 缺列：{', '.join(missing)}（实际：{', '.join(sorted(cols))}）")
    return f"{len(args.get('columns', []))} 列齐全"


def inv_keys_unique(args, base, ctx):
    rows = load_records(args["path"], base)
    key = args["key"]
    seen, dup = set(), []
    for r in rows:
        v = str(_col(r, key))
        if v in seen:
            dup.append(v)
        seen.add(v)
    if dup:
        raise DataError(f"{args['path']} 的 {key} 有重复值：{', '.join(sorted(set(dup))[:5])}")
    return f"{len(rows)} 行键唯一"


def inv_no_null(args, base, ctx):
    rows = load_records(args["path"], base)
    bad = []
    for i, r in enumerate(rows, 1):
        for c in args.get("columns", []):
            if _col(r, c) in (None, ""):
                bad.append(f"第{i}行.{c}")
    if bad:
        raise DataError(f"{args['path']} 有空值：{', '.join(bad[:5])}")
    return "无空值"


def inv_sum_equals(args, base, ctx):
    rows = load_records(args["path"], base)
    total = sum(_num(_col(r, args["column"])) for r in rows)
    want = _num(args["equals"])
    tol = float(args.get("tolerance", 0.01))
    if abs(total - want) > tol:
        raise DataError(
            f"{args['path']}.{args['column']} 合计 {total:.2f} ≠ 期望 {want:.2f}（容差 {tol}）")
    return f"合计 {total:.2f} 相符"


def inv_sums_equal(args, base, ctx):
    left, right = args["left"], args["right"]
    lrows = load_records(left["path"], base)
    rrows = load_records(right["path"], base)
    lt = sum(_num(_col(r, left["column"])) for r in lrows)
    rt = sum(_num(_col(r, right["column"])) for r in rrows)
    tol = float(args.get("tolerance", 0.01))
    if abs(lt - rt) > tol:
        raise DataError(
            f"两侧合计不等：{left['path']}.{left['column']}={lt:.2f} vs "
            f"{right['path']}.{right['column']}={rt:.2f}（差 {lt - rt:.2f}）")
    return f"两侧合计均为 {lt:.2f}"


def inv_equal_counts(args, base, ctx):
    a = load_records(args["left"], base)
    b = load_records(args["right"], base)
    if len(a) != len(b):
        raise DataError(f"行数不等：{args['left']}={len(a)} vs {args['right']}={len(b)}")
    return f"两侧各 {len(a)} 行"


def inv_subset_keys(args, base, ctx):
    left, right = args["left"], args["right"]
    lk = {str(_col(r, left["key"])) for r in load_records(left["path"], base)}
    rk = {str(_col(r, right["key"])) for r in load_records(right["path"], base)}
    miss = sorted(lk - rk)
    if miss:
        raise DataError(
            f"{left['path']} 里有 {len(miss)} 个键不在 {right['path']}：{', '.join(miss[:5])}")
    return f"{len(lk)} 个键全部命中"


def inv_rows_preserved(args, base, ctx):
    """before / after 两组文件的行数守恒（搬数据最常用的一条）。"""
    from .core import expand_globs
    before = sum(len(load_records(p, base)) for p in expand_globs(args["before"], base))
    after = sum(len(load_records(p, base)) for p in expand_globs(args["after"], base))
    if before != after:
        raise DataError(f"行数不守恒：处理前 {before} 行，处理后 {after} 行（差 {before - after}）")
    return f"行数守恒（{before}）"


def inv_field_equals(args, base, ctx):
    """**写后回读**的主力判据：定位到 key==value 那条记录，核它的 field 是不是期望值。

    没有 API 的系统里，这是唯一能落到数据层的"我做对了"的证明——
    比"界面上弹了个成功提示条"硬得多：读错了对象、写没生效、被静默回滚，都会在这里露出来。
    """
    rows = load_records(args["path"], base)
    key, want, field = args["key"], args["value"], args["field"]
    expect = args["equals"]
    hit = None
    for r in rows:
        if str(_col(r, key)).strip() == str(want).strip():
            hit = r
            break
    if hit is None:
        raise DataError(
            f"{args['path']} 里找不到 {key}={want} 的记录——回读没读到目标对象，"
            f"不能认为写成功了（共 {len(rows)} 行）")
    got = _col(hit, field)
    if str(got).strip() != str(expect).strip():
        raise DataError(
            f"回读不符：{key}={want} 的 {field} 实际是 {got!r}，期望 {expect!r}"
            f"（写操作没生效、写错了对象，或被系统回滚）")
    return f"回读一致：{key}={want} 的 {field}={got}"


def inv_field_in(args, base, ctx):
    """回读核对（枚举版）：目标记录的 field 必须落在允许集合里。"""
    rows = load_records(args["path"], base)
    key, want, field = args["key"], args["value"], args["field"]
    allowed = [str(x).strip() for x in (args.get("in") or [])]
    for r in rows:
        if str(_col(r, key)).strip() == str(want).strip():
            got = str(_col(r, field)).strip()
            if got not in allowed:
                raise DataError(
                    f"回读不符：{key}={want} 的 {field}={got!r} 不在允许集合 "
                    f"{allowed} 里")
            return f"回读一致：{key}={want} 的 {field}={got}"
    raise DataError(f"{args['path']} 里找不到 {key}={want} 的记录（回读没读到目标对象）")


def inv_no_duplicate_side_effect(args, base, ctx):
    """写操作只发生了一次：按 key 分组的记录数必须正好是 1。

    重复点提交、resume 时把同一步写了两遍——在 GUI 世界里这是最常见的"假成功"。
    """
    rows = load_records(args["path"], base)
    key = args["key"]
    counts: dict = {}
    for r in rows:
        v = str(_col(r, key)).strip()
        counts[v] = counts.get(v, 0) + 1
    dup = {k: c for k, c in counts.items() if c > 1}
    if dup:
        shown = ", ".join(f"{k}×{c}" for k, c in list(dup.items())[:5])
        raise DataError(f"{args['path']} 的 {key} 出现重复（疑似重复提交）：{shown}")
    return f"{len(counts)} 个键各出现 1 次（没有重复写）"


REGISTRY = {
    "file_exists": inv_file_exists,
    "not_empty": inv_not_empty,
    "row_count_between": inv_row_count_between,
    "columns_present": inv_columns_present,
    "keys_unique": inv_keys_unique,
    "no_null": inv_no_null,
    "sum_equals": inv_sum_equals,
    "sums_equal": inv_sums_equal,
    "equal_counts": inv_equal_counts,
    "subset_keys": inv_subset_keys,
    "rows_preserved": inv_rows_preserved,
    "field_equals": inv_field_equals,
    "field_in": inv_field_in,
    "no_duplicate_side_effect": inv_no_duplicate_side_effect,
}


def check(invariant, base: Path, ctx: dict) -> tuple:
    """跑一条不变量，返回 (是否通过, 说明)。"""
    if isinstance(invariant, str):
        name, args = invariant, {}
    else:
        name = invariant.get("name") or invariant.get("invariant") or ""
        args = {k: v for k, v in invariant.items() if k not in ("name", "invariant", "note")}
    from .core import resolve
    args = resolve(args, ctx)
    fn = REGISTRY.get(name)
    if fn is None:
        raise JobError(f"未知不变量：{name}（可用：{', '.join(sorted(REGISTRY))}）")
    try:
        return True, fn(args, base, ctx)
    except (DataError, EnvError) as exc:
        return False, str(exc)


def unknown_invariants(spec_list) -> list:
    bad = []
    for it in spec_list or []:
        name = it if isinstance(it, str) else (it.get("name") or it.get("invariant"))
        if name not in REGISTRY:
            bad.append(str(name))
    return bad
